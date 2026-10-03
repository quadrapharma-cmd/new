"""Tests for tools/seed on an INVENTED workbook (no company data anywhere in this file).

    python3 -m unittest discover -s tests/seed -v          # or: python3 -m pytest tests/seed
"""
import copy
import glob
import hashlib
import json
import os
import shutil
import subprocess
import sys
import tempfile
import unittest

sys.dont_write_bytecode = True      # keep the repository free of __pycache__
HERE = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HERE, '..', '..'))
TOOLS = os.path.join(REPO, 'tools', 'seed')
sys.path.insert(0, TOOLS)

import make_batches  # noqa: E402
import make_oracle  # noqa: E402
import make_seed  # noqa: E402
import make_synthetic_workbook  # noqa: E402
import parties as P  # noqa: E402
import reader  # noqa: E402
import verify_seed  # noqa: E402
import xref  # noqa: E402
from common import SeedError, assert_outside_repo, write_json  # noqa: E402

OVERRIDES = {'merge': [], 'separate': [], 'kinds': {}, 'owners': {}}


def tree_hash(root):
    h = hashlib.sha256()
    for p in sorted(glob.glob(os.path.join(root, '**', '*.json'), recursive=True)):
        h.update(os.path.relpath(p, root).encode())
        with open(p, 'rb') as fh:
            h.update(fh.read())
    return h.hexdigest()


class SeedBase(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.tmp = tempfile.mkdtemp(prefix='strifa-seed-test-')
        cls.xlsx = os.path.join(cls.tmp, 'synthetic.xlsx')
        make_synthetic_workbook.build(cls.xlsx)
        cls.out = os.path.join(cls.tmp, 'private')
        cls.result = make_seed.build(cls.xlsx, None, copy.deepcopy(OVERRIDES))
        make_seed.write_outputs(cls.result, cls.out)
        cls.docs = cls.result['docs']
        cls.seed_dir = os.path.join(cls.out, 'seed')

    @classmethod
    def tearDownClass(cls):
        shutil.rmtree(cls.tmp, ignore_errors=True)

    def coll(self, name):
        return {k.split('/', 1)[1]: v for k, v in self.docs.items() if k.startswith(name + '/')}

    def entry(self, no):
        return self.docs[f'entries/e{no:06d}']


class TestSchema(SeedBase):
    def test_collections_and_counts(self):
        names = {k.split('/')[0] for k in self.docs}
        self.assertEqual(names, {'meta', 'accounts', 'costCenters', 'sectors', 'parties', 'documents', 'entries', 'fiscalYears',
                                 'openItems', 'assumptions', 'assets', 'templates', 'registers', 'importRuns'})
        self.assertEqual({k for k in self.docs if k.startswith('meta/')}, {'meta/config', 'meta/counters', 'meta/groups'})
        self.assertEqual(len(self.coll('entries')), 20)
        self.assertEqual(len(self.coll('accounts')), 32)
        self.assertEqual(len(self.coll('templates')), 5)

    def test_account_fields(self):
        keys = {'code', 'name', 'notes', 'cls', 'type', 'normal', 'contra', 'contraOf', 'fsLine', 'parent', 'postable', 'active',
                'partyRule', 'everUsed'}
        for code, a in self.coll('accounts').items():
            self.assertEqual(set(a), keys)
            self.assertEqual(a['code'], code)
            self.assertIn(a['normal'], 'DC')
            self.assertIn(a['partyRule'], ('none', 'recommended', 'required'))
        accts = self.coll('accounts')
        self.assertEqual(accts['3120']['contraOf'], '3110')
        self.assertEqual(accts['1192']['contraOf'], '1130')
        self.assertEqual(accts['3211']['fsLine'], 'SETTLEMENT_SHARES')
        self.assertEqual(accts['3150']['fsLine'], 'RETAINED')
        self.assertEqual(accts['3110']['fsLine'], 'CAPITAL')
        self.assertEqual(accts['2310']['partyRule'], 'required')
        self.assertEqual(accts['1340']['partyRule'], 'recommended')

    def test_notes_leave_the_name(self):
        a = self.coll('accounts')
        self.assertEqual(a['2310']['notes'], 'تعهد غير موثق')
        self.assertNotIn('(', a['2310']['name'])
        self.assertTrue(a['5211']['notes'].startswith('مغلق'))
        self.assertNotIn('مغلق', a['5211']['name'])

    def test_active_until_used_and_allow_list(self):
        a = self.coll('accounts')
        self.assertFalse(a['5230']['everUsed'])
        self.assertFalse(a['5230']['active'], 'unused account is inactive until first use')
        for code in ('3150', '2250'):                       # explicit allow-list
            self.assertFalse(a[code]['everUsed'])
            self.assertTrue(a[code]['active'])
        self.assertFalse(a['5211']['active'], 'closed account stays inactive')
        self.assertTrue(a['1340']['everUsed'] and a['1340']['active'])

    def test_entry_and_line_fields(self):
        ekeys = {'no', 'date', 'fy', 'status', 'desc', 'docIds', 'cc', 'sector', 'source', 'isLegacy', 'legacyRow', 'version',
                 'voidReason', 'reversalOf', 'createdBy', 'createdAt', 'postedBy', 'postedAt', 'lines', 'history'}
        lkeys = {'n', 'acct', 'dr', 'cr', 'memo', 'partyId', 'docIds', 'cc', 'sector', 'valueDate', 'needsReview', 'reviewReason',
                 'links', 'legacy'}
        for eid, e in self.coll('entries').items():
            self.assertEqual(set(e), ekeys)
            self.assertEqual(eid, f'e{e["no"]:06d}')
            self.assertEqual((e['status'], e['isLegacy'], e['source'], e['version']), ('posted', True, 'legacy', 1))
            self.assertEqual(e['fy'], int(e['date'][:4]))
            for l in e['lines']:
                self.assertEqual(set(l), lkeys)
                self.assertEqual(set(l['legacy']), {'partyText', 'docText', 'row'})
                self.assertTrue((l['dr'] > 0) != (l['cr'] > 0))
            self.assertEqual([l['n'] for l in e['lines']], list(range(1, len(e['lines']) + 1)))

    def test_party_document_dimension_fields(self):
        for p in self.coll('parties').values():
            self.assertTrue({'id', 'name', 'kind', 'aliases', 'roles', 'taxId', 'active', 'mergedInto', 'defaultAccount',
                             'legacyTexts'} <= set(p))
            self.assertIn(p['kind'], P.PARTY_KINDS)
        for d in self.coll('documents').values():
            self.assertEqual(set(d), {'id', 'ref', 'type', 'date', 'grade', 'status', 'note', 'files', 'legacyUseCount'})
            self.assertEqual((d['status'], d['files']), ('expected', []))
        for c in self.coll('costCenters').values():
            self.assertEqual(set(c), {'id', 'name', 'kind', 'active'})
        for s in self.coll('sectors').values():
            self.assertEqual(set(s), {'id', 'name', 'active'})

    def test_meta_docs(self):
        cfg = self.docs['meta/config']
        self.assertEqual(cfg['schemaVersion'], 1)
        self.assertEqual(cfg['company']['incorporated'], '2019-03-01')
        self.assertEqual(cfg['company']['formerName'], 'شركة الأمل للتجارة')
        b = cfg['baseline']
        self.assertEqual(len(b['sha256']), 64)
        self.assertEqual(b['counts']['entries'], 20)
        self.assertEqual(self.docs['meta/counters']['nextEntryNo'], 21)
        self.assertIn('accrualAccountOf', cfg['rules'])

    def test_sectors_and_cost_centers(self):
        names = {s['name'] for s in self.coll('sectors').values()}
        self.assertIn('خارج ق.72', names)
        self.assertEqual(self.docs['meta/config']['defaults']['sector'],
                         next(s['id'] for s in self.coll('sectors').values() if s['name'] == 'مشترك'))
        self.assertEqual({c['name']: c['kind'] for c in self.coll('costCenters').values()}, {'عام': 'own', 'المصنع': 'own'})

    def test_fiscal_years_open_items_assumptions(self):
        fy = self.coll('fiscalYears')
        self.assertEqual({k: v['state'] for k, v in fy.items()}, {'2019': 'closed_reserved', '2020': 'closed_reserved', '2021': 'open'})
        self.assertEqual(fy['2019']['legalStart'], '2019-03-01')
        self.assertEqual(fy['2019']['reservation']['lastUpdate'], '2026-02-05')
        self.assertEqual([n['text'] for n in fy['2019']['neededDocs']], ['كشف حساب البنك من التأسيس', 'إيصالات مصروفات التأسيس'])
        self.assertIsNotNone(fy['2019']['snapshot'])
        self.assertIsNone(fy['2021']['snapshot'])
        oi = self.coll('openItems')
        self.assertEqual({k: v['status'] for k, v in oi.items()},
                         {'1': 'open', '2': 'inquiry', '3': 'partial', '4': 'closed', '5': 'open', '6': 'open'})
        self.assertEqual((oi['3']['year'], oi['3']['from'], oi['3']['to']), ('2019-2020', 2019, 2020))
        self.assertEqual((oi['5']['year'], oi['5']['from'], oi['5']['to']), ('2019+', 2019, None))
        self.assertEqual((oi['4']['year'], oi['4']['from'], oi['4']['to']), (2020, 2020, 2020))
        self.assertTrue(all(a['status'] == 'pending' for a in self.coll('assumptions').values()))


class TestBalanceAndHierarchy(SeedBase):
    def test_every_entry_balanced(self):
        for e in self.coll('entries').values():
            dr = sum(round(l['dr'] * 100) for l in e['lines'])
            cr = sum(round(l['cr'] * 100) for l in e['lines'])
            self.assertEqual(dr, cr, e['no'])
            self.assertGreaterEqual(len(e['lines']), 2)
        self.assertTrue(self.result['report']['totalDebitCents'] == self.result['report']['totalCreditCents'])

    def test_multi_line_entry_and_row_order(self):
        e = self.entry(17)                                  # the 2020 depreciation entry
        self.assertEqual(len(e['lines']), 9)
        rows = [l['legacy']['row'] for l in e['lines']]
        self.assertEqual(rows, sorted(rows))
        self.assertEqual(self.entry(9)['legacyRow'], 2)     # typed first although numbered 9

    def test_hierarchy_derived_from_codes(self):
        groups = self.docs['meta/groups']['groups']
        l1 = {g['code'] for g in groups if g['level'] == 1}
        l2 = {g['code'] for g in groups if g['level'] == 2}
        self.assertEqual(l1, set('123456'))
        self.assertEqual(l2, {a['code'][:2] for a in self.coll('accounts').values()})
        for a in self.coll('accounts').values():
            self.assertEqual(a['parent'], a['code'][:2])
            self.assertTrue(a['postable'])
            self.assertIn(a['parent'], l2)
        names = {g['code']: g['name'] for g in groups}
        self.assertEqual(names['1'], 'الأصول')
        self.assertEqual(names['5'], 'التكاليف والمصروفات')

    def test_header_dates_and_value_date(self):
        e = self.entry(16)
        self.assertEqual(e['date'], '2020-12-31')
        self.assertEqual([l['valueDate'] for l in e['lines']], [None, None, '2020-03-04'])
        self.assertTrue(all(l['valueDate'] is None for l in self.entry(9)['lines']))

    def test_header_dimensions_are_the_most_common_value(self):
        e = self.entry(18)                                  # one entry carrying the other cost centre on every line
        factory = next(c['id'] for c in self.coll('costCenters').values() if c['name'] == 'المصنع')
        self.assertEqual(e['cc'], factory)
        self.assertTrue(all(l['cc'] is None for l in e['lines']))
        e20 = self.entry(20)
        s72 = next(s['id'] for s in self.coll('sectors').values() if s['name'] == 'داخل ق.72')
        self.assertEqual(e20['sector'], s72)


class TestCrossReferences(unittest.TestCase):
    def test_parse_links(self):
        cases = {
            'قيد — افتراض 33': ([33], []),
            'افتراض رقم 7': ([7], []),
            'بند مفتوح 48': ([], [48]),
            'قيد الاستفسار (بند 57)': ([], [57]),
            'لتجمع قيد الاستفسار 76': ([], [76]),
            'فرق — استفسار 60)': ([], [60]),
            'افتراض 38 — استفسار بند 56': ([38], [56]),
            'افتراض ٣٨': ([38], []),
            '(البند 2-7) ألزمت': ([], []),                # contract clause cited with the article
            'بند 4-5 من العقد': ([], []),                 # clause number N-M
            '41 بنداً بإجمالي': ([], []),                 # a count of items
            'قيد الاستفسار': ([], []),                    # the tag alone is not a link
        }
        for text, want in cases.items():
            self.assertEqual(xref.parse_links(text), want, text)
        self.assertTrue(xref.has_inquiry_tag('خطأ — استفسار: لم يُعثر'))
        self.assertFalse(xref.has_inquiry_tag('نص عادي'))


class TestLinksInSeed(SeedBase):
    def test_links_set_and_flagged(self):
        e16 = self.entry(16)
        line = e16['lines'][2]
        self.assertEqual(line['links'], [{'t': 'a', 'n': 2}, {'t': 'o', 'n': 4}])
        self.assertTrue(line['needsReview'])
        self.assertIn('افتراض 2', line['reviewReason'])
        self.assertEqual(e16['lines'][0]['links'], [], 'the contract-clause look-alike must not become a link')
        self.assertFalse(e16['lines'][0]['needsReview'])

    def test_inquiry_and_register_links(self):
        e5 = self.entry(5)
        self.assertIn('قيد الاستفسار', e5['lines'][0]['reviewReason'])
        oi, asm = self.coll('openItems'), self.coll('assumptions')
        self.assertEqual(oi['4']['linkedEntries'], [16])
        self.assertEqual(oi['2']['linkedEntries'], [3])
        self.assertEqual(asm['1']['linkedEntries'], [6])
        self.assertEqual(asm['3']['linkedEntries'], [9])
        for e in self.coll('entries').values():
            for l in e['lines']:
                for k in l['links']:
                    reg = oi if k['t'] == 'o' else asm
                    self.assertIn(str(k['n']), reg)
                    self.assertIn(e['no'], reg[str(k['n'])]['linkedEntries'])

    def test_duplicates_and_pre_incorporation_flagged(self):
        for no in (12, 13):
            self.assertTrue(all(l['needsReview'] and 'تكرار محتمل' in l['reviewReason'] for l in self.entry(no)['lines']))
            self.assertEqual(self.entry(no)['no'], no)       # never merged
        e9 = self.entry(9)
        self.assertTrue(all(l['needsReview'] and 'يسبق تأسيس' in l['reviewReason'] for l in e9['lines']))
        self.assertIn({'t': 'o', 'n': 6}, e9['lines'][0]['links'])   # the open item that cites that date

    def test_party_required_missing_flagged(self):
        l = self.entry(5)['lines'][0]                        # 1290 requires a party, 'غير محدد' gives none
        self.assertIsNone(l['partyId'])
        self.assertIn('يشترطه', l['reviewReason'])


class TestParties(SeedBase):
    def party(self, name):
        return next(p for p in self.coll('parties').values() if name in [p['name']] + p['aliases'])

    def test_placeholders_become_null_with_legacy_text(self):
        for no, idx, text in ((3, 0, '—'), (4, 0, 'سباك'), (4, 1, 'عام'), (5, 0, 'غير محدد'), (8, 0, '')):
            l = self.entry(no)['lines'][idx]
            self.assertEqual(l['legacy']['partyText'], text)
            if text != '—' or idx != 1:
                pass
        self.assertIsNone(self.entry(4)['lines'][0]['partyId'])
        self.assertIsNone(self.entry(3)['lines'][0]['partyId'])         # bank line with a placeholder
        self.assertIsNone(self.entry(8)['lines'][0]['partyId'])
        names = [n for p in self.coll('parties').values() for n in [p['name']] + p['aliases']]
        for ph in ('—', 'عام', 'غير محدد', 'مورد', 'سباك', 'مورد الأسمنت'):
            self.assertNotIn(ph, names)

    def test_funding_account_takes_owner_from_account(self):
        owner = self.party('د/ زهرة المقدسي')
        l = self.entry(3)['lines'][1]                                    # 2310 line typed with '—'
        self.assertEqual(l['legacy']['partyText'], '—')
        self.assertEqual(l['partyId'], owner['id'])
        self.assertEqual(owner['defaultAccount'], '2310')
        self.assertEqual(owner['accrualAccount'], '2311')
        self.assertEqual(owner['boundAccounts'], ['2310', '2311', '2322'])
        self.assertEqual(self.docs['meta/config']['rules']['accrualAccountOf'], {'2310': '2311'})
        self.assertEqual(self.result['report']['partyOwnerSubstitutedLines'], 1)

    def test_safe_clusters_merged(self):
        self.assertEqual(self.party('زهرة المقدسي (قيد الاستفسار)')['id'], self.party('د/ زهرة المقدسي')['id'])
        self.assertEqual(self.party('بنك الاختبار الوطني')['id'], self.party('الاختبار الوطني')['id'])
        self.assertEqual(self.party('سامية الراوي')['id'], self.party('سامية علي الراوي')['id'])
        merged = self.party('سامية الراوي')
        self.assertEqual(merged['name'], 'سامية علي الراوي')              # the fuller spelling
        self.assertEqual(sorted(merged['legacyTexts']), ['سامية الراوي', 'سامية علي الراوي'])

    def test_lookalike_kept_separate_and_flagged(self):
        a, b = self.party('منير الحلبي'), self.party('أ/ منير')
        self.assertNotEqual(a['id'], b['id'])
        with open(os.path.join(self.out, 'review', 'party_alias_review.csv'), encoding='utf-8-sig') as fh:
            csv = fh.read().splitlines()
        self.assertEqual(csv[0], 'raw,lines,years,top_accounts,proposed_party,kind,confidence,identity_sensitive')
        rows = {r.split(',')[0]: r.split(',') for r in csv[1:]}
        self.assertEqual(rows['أ/ منير'][-1], '1')
        self.assertEqual(rows['منير الحلبي'][-1], '1')
        self.assertEqual(rows['د/ زهرة المقدسي'][-1], '0')

    def test_groups_and_kinds(self):
        self.assertEqual(self.party('المساهمون')['kind'], 'group')
        self.assertEqual(self.party('عملاء')['kind'], 'group')
        self.assertEqual(self.party('منير الحلبي')['kind'], 'shareholder')
        # a share account first posted in the incorporation year -> the role starts at incorporation
        self.assertEqual(self.party('منير الحلبي')['roles'], [{'role': 'shareholder', 'from': '2019-03-01', 'to': None}])
        self.assertEqual(self.party('د/ زهرة المقدسي')['kind'], 'financier')
        self.assertEqual(self.party('بنك الاختبار الوطني')['kind'], 'bank')

    def test_overrides_force_merge_and_separate(self):
        ov = {'merge': [['أ/ منير', 'منير الحلبي']], 'separate': [], 'kinds': {'مصنع الفجر': 'professional'}, 'owners': {}}
        r = make_seed.build(self.xlsx, None, ov)
        parties = [v for k, v in r['docs'].items() if k.startswith('parties/')]
        merged = [p for p in parties if 'أ/ منير' in p['legacyTexts']][0]
        self.assertIn('منير الحلبي', merged['legacyTexts'])
        self.assertEqual(next(p for p in parties if p['name'] == 'مصنع الفجر')['kind'], 'professional')
        ov2 = {'merge': [['أ/ منير', 'منير الحلبي']], 'separate': [['أ/ منير', 'منير الحلبي']], 'kinds': {}, 'owners': {}}
        r2 = make_seed.build(self.xlsx, None, ov2)
        parties2 = [v for k, v in r2['docs'].items() if k.startswith('parties/')]
        self.assertFalse(any('أ/ منير' in p['legacyTexts'] and 'منير الحلبي' in p['legacyTexts'] for p in parties2))

    def test_classification_and_keys(self):
        self.assertEqual([P.classify(x) for x in ('عام', '—', 'سباك', 'مورد الأسمنت', 'المساهمون', 'عملاء نقدي', 'اسم ما')],
                         ['placeholder', 'placeholder', 'trade', 'trade', 'group', 'group', 'party'])
        self.assertEqual(P.key_tokens('د/ عبد الفتاح'), P.key_tokens('الدكتور عبدالفتاح'))
        self.assertEqual(P.key_tokens('بنك الاختبار الوطني'), P.key_tokens('الاختبار الوطني'))


class TestDocuments(SeedBase):
    def test_one_to_one_no_plus_split(self):
        docs = self.coll('documents')
        refs = [d['ref'] for d in docs.values()]
        self.assertEqual(len(refs), len(set(refs)))
        self.assertIn('شهادة تأسيس + استمارة سجل', refs)            # '+' kept inside one document
        raw = {l['doc'] for l in reader.read_journal(reader.open_workbook(self.xlsx))['lines']}
        self.assertEqual(set(refs), raw)
        use = {d['ref']: d['legacyUseCount'] for d in docs.values()}
        self.assertEqual(use['شهادة تأسيس + استمارة سجل'], 3)
        self.assertEqual(sum(use.values()), len([l for l in reader.read_journal(reader.open_workbook(self.xlsx))['lines'] if l['doc']]))
        self.assertEqual(sorted(docs), [f'D{i:04d}' for i in range(1, len(docs) + 1)])

    def test_type_and_grade(self):
        by = {d['ref']: d for d in self.coll('documents').values()}
        self.assertEqual(by['كشف بنك الاختبار 7788 — شيك 501 + كشف حسابات الشركاء']['type'], 'كشف بنك')
        self.assertEqual(by['كشف بنك الاختبار 7788 — شيك 501 + كشف حسابات الشركاء']['grade'], 'A')
        self.assertEqual(by['دفتر الخزينة (بخط اليد)']['grade'], 'B')
        self.assertEqual(by['إفادة الإدارة — تمويل مبكر']['grade'], 'C')
        self.assertIsNone(by['جدول المرتبات']['grade'])


class TestAssets(SeedBase):
    def test_register(self):
        a = sorted(self.coll('assets').values(), key=lambda x: x['id'])
        by = {x['name']: x for x in a}
        self.assertEqual(by['آلات ومعدات']['convention'], 'half_year')
        self.assertEqual(by['آلات ومعدات']['inServiceDate'], '2019-07-02')
        self.assertEqual(by['آلات ومعدات']['accumAcct'], '1192')
        self.assertEqual(by['أرض']['convention'], 'none')
        self.assertEqual((by['كاميرات المراقبة']['convention'], by['كاميرات المراقبة']['inServiceDate']), ('full_next_year', '2019-12-31'))
        self.assertEqual((by['جهاز البصمة']['convention'], by['جهاز البصمة']['inServiceDate']), ('day_count', '2020-03-07'))
        self.assertEqual((by['طابعة']['convention'], by['طابعة']['sourceEntryNo']), ('half_year', 15))
        self.assertEqual(by['كاميرات المراقبة']['sourceEntryNo'], 7)
        self.assertEqual(by['كاميرات المراقبة']['rate'], 0.2)
        cip = by['مشروعات تحت التنفيذ — 2020']
        self.assertEqual((cip['convention'], cip['cost']), ('none', 10000))
        ret = next(x for x in a if x['cost'] < 0)
        self.assertEqual((ret['acct'], ret['cost'], ret['convention']), ('1130', -500, 'none'))
        for x in a:
            self.assertIn(x['convention'], ('half_year', 'day_count', 'full_next_year', 'full_year', 'none'))
            self.assertEqual(x['sourceEntryNo'] in {e['no'] for e in self.coll('entries').values()}, True)

    def test_replay_matches_posted_depreciation(self):
        oracle, ok, _, _ = make_oracle.build(self.xlsx, self.seed_dir)
        self.assertTrue(ok)
        chk = verify_seed.verify(self.seed_dir, oracle)
        self.assertEqual(chk.mismatches, 0, dict(chk.failed))
        self.assertGreater(chk.total, 500)

    def test_charge_math_half_up(self):
        from assets_infer import charge_cents
        self.assertEqual(charge_cents(4111300, 0.05, 'half_year', 2019, '2019-07-02'), 102783)            # 102,782.5 -> half-up
        self.assertEqual(charge_cents(320000, 0.2, 'day_count', 2023, '2023-01-11'), 62247)               # 3,200 x 20% x 355/365
        self.assertEqual(charge_cents(100, 0.5, 'half_year', 2020, '2020-01-01'), 25)
        self.assertEqual(charge_cents(100, 0.2, 'none', 2020, '2020-01-01'), 0)


class TestRegisters(SeedBase):
    def test_registers_present_values_only(self):
        regs = {k.split('/')[1]: v for k, v in self.docs.items() if k.startswith('registers/')}
        self.assertEqual(set(regs), {'settlementCreditors', 'contingentClaims', 'financierFunding', 'legalExpenses',
                                     'payrollExposure', 'accountantReconciliation', 'fundingAllocation', 'fundingEvidence',
                                     'cashSourcesUses', 'workbookNotes'})
        for name, r in regs.items():
            self.assertIn('items', r)
            self.assertIn('source', r)
        self.assertEqual(len(regs['legalExpenses']['items']), 3)
        self.assertEqual(regs['legalExpenses']['total'], 1100)
        self.assertEqual(regs['settlementCreditors']['items'][0]['acct'], '3211')
        self.assertEqual(regs['settlementCreditors']['items'][0]['legacyCode'], '2351')

    def test_evidence_grades_and_status_normalised(self):
        items = {i['n']: i for i in self.docs['registers/fundingEvidence']['items']}
        self.assertEqual([items[i]['grade'] for i in (1, 2, 3, 4)], ['A', 'B', 'C', 'B'])
        self.assertEqual([items[i]['status'] for i in (1, 2, 3, 4)], ['included', 'included', 'after_close', 'excluded'])
        self.assertEqual(reader.normalise_grade('(ب) + تطابق مع المحاسب'), ('B', '+ تطابق مع المحاسب'))
        self.assertEqual(reader.normalise_grade('(أ) — الجزء المُنفَق'), ('A', '— الجزء المُنفَق'))
        self.assertEqual(reader.normalise_grade('(ج) — لا عقد'), ('C', '— لا عقد'))
        self.assertEqual(reader.normalise_grade('بلا قوس')[0], None)

    def test_allocation_formula_literals(self):
        row = self.docs['registers/fundingAllocation']['items'][0]
        self.assertEqual(row['inputFormulas']['inj2022'], {'components': [2000, 3000]})
        self.assertEqual(len(self.docs['registers/fundingAllocation']['deposits']), 2)

    def test_docs_stay_below_200kib(self):
        for path in glob.glob(os.path.join(self.seed_dir, '*', '*.json')):
            self.assertLessEqual(os.path.getsize(path), 200 * 1024, path)


class TestDeterminism(SeedBase):
    def test_build_twice_is_identical(self):
        again = make_seed.build(self.xlsx, None, copy.deepcopy(OVERRIDES))
        self.assertEqual(json.dumps(again['docs'], sort_keys=True), json.dumps(self.docs, sort_keys=True))

    def test_rerun_rewrites_the_same_bytes(self):
        before = tree_hash(self.seed_dir)
        make_seed.write_outputs(self.result, self.out)
        self.assertEqual(tree_hash(self.seed_dir), before)

    def test_independent_of_hash_seed(self):
        hashes = []
        for seed in ('1', '4242'):
            out = os.path.join(self.tmp, f'hs{seed}')
            env = dict(os.environ, PYTHONHASHSEED=seed, STRIFA_PRIVATE_DIR=out)
            subprocess.run([sys.executable, os.path.join(TOOLS, 'make_seed.py'), '--xlsx', self.xlsx, '--overrides',
                            os.path.join(self.tmp, 'none.json')], check=True, capture_output=True, env=env)
            hashes.append(tree_hash(os.path.join(out, 'seed')))
        self.assertEqual(hashes[0], hashes[1])
        self.assertEqual(hashes[0], tree_hash(self.seed_dir))

    def test_stale_files_are_removed(self):
        stale = os.path.join(self.seed_dir, 'parties', 'p9999.json')
        write_json(stale, {'id': 'p9999'})
        make_seed.write_outputs(self.result, self.out)
        self.assertFalse(os.path.exists(stale))


class TestOracleAndBatches(SeedBase):
    def test_oracle_self_check_passes(self):
        oracle, ok, n_checks, n_bad = make_oracle.build(self.xlsx, self.seed_dir)
        self.assertTrue(ok, oracle['selfCheck'])
        self.assertEqual(n_bad, 0)
        self.assertGreater(n_checks, 100)
        self.assertEqual(sorted(oracle['years']), ['2019', '2020'])
        self.assertEqual(oracle['counts']['entries'], 20)
        self.assertEqual(oracle['counts']['unbalancedEntries'], 0)
        self.assertEqual(oracle['counts']['linesCitingAssumption'], 3)
        self.assertEqual(oracle['counts']['linesCitingOpenItem'], 3)
        self.assertEqual(oracle['counts']['accountsUnused'], 32 - oracle['counts']['accountsUsed'])
        y20 = oracle['years']['2020']
        self.assertEqual(y20['balanceCheck'], 0)
        self.assertEqual(oracle['years']['2019']['openingGross'], {'debit': 0, 'credit': 0})
        self.assertEqual(y20['openingGross']['debit'], oracle['years']['2019']['debit'])
        tb = oracle['reports']['2020']['trialBalance']
        self.assertAlmostEqual(tb['debit'], tb['credit'])
        self.assertEqual(oracle['accounts']['3110']['2019']['closeCr'], 1000000)

    def test_oracle_detects_a_wrong_cached_figure(self):
        import openpyxl
        wb = openpyxl.load_workbook(self.xlsx)
        wb.worksheets[5]['B4'] = 123456789                      # corrupt the cached revenue figure
        bad = os.path.join(self.tmp, 'bad.xlsx')
        wb.save(bad)
        _, ok, _, n_bad = make_oracle.build(bad, None)
        self.assertFalse(ok)
        self.assertGreater(n_bad, 0)

    def test_batches(self):
        out = os.path.join(self.tmp, 'batch-out')
        shutil.copytree(self.seed_dir, os.path.join(out, 'seed'))
        self.assertEqual(make_batches.main(['--out', out, '--max', '50']), 0)
        with open(os.path.join(out, 'batches', 'index.json'), encoding='utf-8') as fh:
            index = json.load(fh)
        self.assertEqual(index['docCount'], len(self.docs))
        files = sorted(glob.glob(os.path.join(out, 'batches', 'batch_*.json')))
        self.assertEqual(len(files), index['batchCount'])
        writes = []
        for f in files:
            with open(f, encoding='utf-8') as fh:
                w = json.load(fh)['writes']
            self.assertLessEqual(len(w), 50)
            writes += w
        self.assertEqual(len(writes), len(self.docs))
        self.assertTrue(all(w['op'] == 'set' and os.path.isabs(w['file_path']) and os.path.exists(w['file_path']) for w in writes))
        colls = [w['collection'] for w in writes]
        self.assertEqual([w['doc_id'] for w in writes[:3]], ['config', 'counters', 'groups'])
        first = {c: colls.index(c) for c in set(colls)}
        last = {c: len(colls) - 1 - colls[::-1].index(c) for c in set(colls)}
        self.assertEqual(colls[0], 'meta')
        self.assertLess(last['meta'], first['accounts'])
        self.assertLess(last['accounts'], first['entries'])
        self.assertLess(last['entries'], first['registers'])
        self.assertEqual(len({(w['collection'], w['doc_id']) for w in writes}), len(writes))
        ids = [w['doc_id'] for w in writes if w['collection'] == 'openItems']
        self.assertEqual(ids, [str(i) for i in range(1, 7)])


class TestGuards(unittest.TestCase):
    def test_refuses_to_write_inside_the_repo(self):
        with self.assertRaises(SeedError):
            assert_outside_repo(os.path.join(REPO, 'seed-output'))
        with self.assertRaises(SeedError):
            assert_outside_repo(REPO)
        assert_outside_repo(tempfile.gettempdir())

    def test_document_size_cap(self):
        with tempfile.TemporaryDirectory() as t:
            with self.assertRaises(SeedError):
                write_json(os.path.join(t, 'big.json'), {'x': 'ا' * 120000})
            self.assertEqual(write_json(os.path.join(t, 'ok.json'), {'x': 1}), 8)

    def test_bad_workbook_stops_the_run(self):
        import openpyxl
        wb = openpyxl.Workbook()
        p = os.path.join(tempfile.mkdtemp(), 'x.xlsx')
        wb.save(p)
        with self.assertRaises(SeedError):
            make_seed.build(p)

    def test_unbalanced_entry_is_rejected(self):
        import openpyxl
        with tempfile.TemporaryDirectory() as t:
            src = os.path.join(t, 's.xlsx')
            make_synthetic_workbook.build(src)
            wb = openpyxl.load_workbook(src)
            wb.worksheets[2]['F2'] = 49999                         # break the balance of the first entry
            bad = os.path.join(t, 'b.xlsx')
            wb.save(bad)
            with self.assertRaises(SeedError):
                make_seed.build(bad)

    def test_templates_match_the_engine(self):
        engine = os.path.join(REPO, 'src', 'engine', 'templates.js')
        node = shutil.which('node')
        if not (node and os.path.exists(engine)):
            self.skipTest('node or src/engine/templates.js not available')
        code = ("import { defaultTemplates } from '%s'; process.stdout.write(JSON.stringify(defaultTemplates()));" % engine)
        res = subprocess.run([node, '--input-type=module', '-e', code], capture_output=True, text=True, timeout=60)
        if res.returncode != 0:
            self.skipTest('engine module could not be loaded: ' + res.stderr[:200])
        import templates_seed
        self.assertEqual(json.loads(res.stdout), templates_seed.templates(),
                         'templates_seed.py drifted from defaultTemplates() in src/engine/templates.js')


if __name__ == '__main__':
    unittest.main()

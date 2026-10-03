#!/usr/bin/env python3
"""Independent acceptance numbers ("oracle") for the migrated books.

    python3 tools/seed/make_oracle.py [--xlsx PATH] [--out DIR] [--seed DIR]

Computed WITHOUT reading the seed JSON: the journal is replayed with pandas (integer cents) straight from the workbook, and
the workbook's own cached results (sheets 3-9, 15-20) are read next to it. A consistency pass compares the replay with the
cached 2025 statements and prints PASS/FAIL with counts only. Output: <out>/oracle.json (private; never inside the repo).
Only `counts.partiesAfterMerge` / `counts.assetsRegister` are read from the seed when it exists (marked seedDerived).
"""
import argparse
import json
import os
import re
import sys
from collections import Counter, defaultdict

import pandas as pd

sys.dont_write_bytecode = True      # keep the repository free of __pycache__
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from common import (SeedError, arabic_index, assert_outside_repo, clean, dumps, eprint, from_cents, private_dir,  # noqa: E402
                    sha256_file, to_western, xlsx_path)

PL_CLASSES = ('إيرادات', 'إيرادات مدين', 'تكلفة مبيعات', 'مصروفات', 'مصروفات غير واجبة الخصم')
CASH_PREFIX = '13'            # group 13 = cash and banks in this chart
FUNDING_ACCOUNTS = ('2310', '2311', '2320', '2321', '2322', '2324', '2330', '2340')
PLACEHOLDER_TEXTS = {'', '—', 'عام', 'غير محدد', 'مورد', 'موردون'}


def cents_of(series):
    s = pd.to_numeric(series, errors='coerce').fillna(0.0)
    c = (s * 100).round().astype('int64')
    if ((s * 100 - c).abs() > 1e-6).any():
        raise SeedError('journal amount with more than two decimals')
    return c


def load_frames(path):
    j = pd.read_excel(path, sheet_name=2, header=0, engine='openpyxl')
    need = ['التاريخ', 'رقم القيد', 'كود الحساب', 'مدين', 'دائن', 'البيان', 'رقم المستند', 'مركز التكلفة', 'الطرف',
            'القطاع القانوني']
    for c in need:
        if c not in j.columns:
            raise SeedError('journal column missing: ' + c)
    j = j[j['كود الحساب'].notna() & (j['كود الحساب'].astype(str).str.strip() != '')].copy()
    j['acct'] = j['كود الحساب'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
    j['date'] = pd.to_datetime(j['التاريخ'])
    j['year'] = j['date'].dt.year
    j['no'] = j['رقم القيد'].astype(float).astype('int64')
    j['dr'] = cents_of(j['مدين'])
    j['cr'] = cents_of(j['دائن'])
    c = pd.read_excel(path, sheet_name=1, header=0, engine='openpyxl')
    c = c[c['الكود'].notna() & c['التصنيف'].notna()].copy()
    c['code'] = c['الكود'].astype(str).str.strip().str.replace(r'\.0$', '', regex=True)
    c = c[c['code'].str.fullmatch(r'[1-9]\d{3}')]
    return j, c[['code', 'اسم الحساب', 'التصنيف']].rename(columns={'اسم الحساب': 'name', 'التصنيف': 'cls'})


def money(c):
    return from_cents(int(c))


def split_dc(net):
    return (money(max(net, 0)), money(max(-net, 0)))


def main(argv=None):
    ap = argparse.ArgumentParser()
    ap.add_argument('--xlsx')
    ap.add_argument('--out')
    ap.add_argument('--seed')
    args = ap.parse_args(argv)
    out_dir = args.out or private_dir()
    assert_outside_repo(out_dir)
    path = xlsx_path(args.xlsx)
    seed_dir = args.seed or os.path.join(out_dir, 'seed')
    try:
        oracle, ok, n_checks, n_bad = build(path, seed_dir)
    except SeedError as e:
        eprint(f'make_oracle: {e}')
        return 2
    os.makedirs(out_dir, exist_ok=True)
    target = os.path.join(out_dir, 'oracle.json')
    with open(target, 'w', encoding='utf-8') as fh:
        fh.write(dumps(oracle) + '\n')
    print(f'make_oracle: wrote oracle ({len(oracle["years"])} years, {len(oracle["accounts"])} accounts)')
    print(f'self-check vs workbook cached values: {"PASS" if ok else "FAIL"} (checks={n_checks}, mismatches={n_bad})')
    return 0 if ok else 1


def cached_by_label(ws, col_label=1, col_value=2, stop_at=None):
    out = {}
    for r in range(1, ws.max_row + 1):
        lab = ws.cell(r, col_label).value
        if isinstance(lab, str):
            out.setdefault(clean(lab), ws.cell(r, col_value).value)
    return out


def num_or_none(v):
    return None if v is None or isinstance(v, str) else v


def cents_cached(v):
    if v is None or isinstance(v, str):
        return None
    return int(round(float(v) * 100))


def build(path, seed_dir=None):
    import openpyxl
    j, chart = load_frames(path)
    wbv = openpyxl.load_workbook(path, data_only=True)
    sheets = wbv.worksheets
    if len(sheets) != 21:
        raise SeedError('unexpected sheet count')
    for i, ws in enumerate(sheets):
        if not ws.title.strip().startswith(arabic_index(i) + '.'):
            raise SeedError(f'sheet {i} title mismatch')
    cls_of = dict(zip(chart['code'], chart['cls']))
    codes = list(chart['code'])
    unknown = sorted(set(j['acct']) - set(codes))
    if unknown:
        raise SeedError('journal uses codes missing from the chart')
    years = sorted(int(y) for y in j['year'].unique())
    checks, bad, failed = [0], [0], Counter()

    def check(cond):
        checks[0] += 1
        if not cond:
            bad[0] += 1
            failed[sys._getframe(1).f_lineno] += 1       # source line of the failing check (no data in the message)
        return bool(cond)

    # ------------------------------------------------------------------ replay
    j['net'] = j['dr'] - j['cr']
    per = j.groupby(['year', 'acct']).agg(mdr=('dr', 'sum'), mcr=('cr', 'sum')).reset_index()
    mov = {(int(r.year), r.acct): (int(r.mdr), int(r.mcr)) for r in per.itertuples()}
    oracle = {'meta': {'workbookSha256': sha256_file(path), 'generator': 'tools/seed/make_oracle.py',
                       'note': 'independent replay of the journal + workbook cached values; amounts in EGP',
                       'tolerance': {'money': 0.005, 'ratio': 1e-9}},
              'years': {}, 'accounts': {}}
    cum = defaultdict(int)          # acct -> cumulative net (cents) after the last processed year
    gdr, gcr = defaultdict(int), defaultdict(int)   # acct -> cumulative GROSS debits / credits (what sheet 3 shows as opening)
    prior_net_result = 0
    snapshots = {}
    for y in years:
        opening = dict(cum)
        open_gdr, open_gcr = dict(gdr), dict(gcr)
        yr_dr = yr_cr = 0
        for a in codes:
            d, c = mov.get((y, a), (0, 0))
            yr_dr += d
            yr_cr += c
            cum[a] += d - c
            gdr[a] += d
            gcr[a] += c
        net_result = -sum(mov.get((y, a), (0, 0))[0] - mov.get((y, a), (0, 0))[1] for a in codes if cls_of[a] in PL_CLASSES)
        assets = sum(cum[a] for a in codes if cls_of[a] == 'أصول')
        accum = -sum(cum[a] for a in codes if cls_of[a] == 'مجمع إهلاك')
        liab = -sum(cum[a] for a in codes if cls_of[a] == 'التزامات')
        susp = -sum(cum[a] for a in codes if cls_of[a] == 'وسيط')
        # equity: credit-normal 'حقوق ملكية' adds, the debit contra 'حقوق ملكية مدين' subtracts, plus every P&L result
        cap_called = sum(cum[a] for a in codes if cls_of[a] == 'حقوق ملكية مدين')
        eq_credit = -sum(cum[a] for a in codes if cls_of[a] == 'حقوق ملكية')
        cum_pl = -sum(cum[a] for a in codes if cls_of[a] in PL_CLASSES)     # all P&L through this year-end (cumulative)
        equity = eq_credit - cap_called + cum_pl
        net_assets = assets - accum
        tb_dr = sum(max(v, 0) for v in cum.values())
        tb_cr = sum(max(-v, 0) for v in cum.values())
        oracle['years'][str(y)] = {
            'debit': money(yr_dr), 'credit': money(yr_cr), 'netResult': money(net_result),
            'assetsGross': money(assets), 'accumDep': money(accum), 'netAssets': money(net_assets),
            'liabilities': money(liab), 'suspense': money(susp), 'equity': money(equity),
            'equityParts': {'capitalAndReserves': money(eq_credit), 'capitalCalledDebit': money(cap_called),
                            'priorResults': money(cum_pl - net_result), 'currentResult': money(net_result)},
            'balanceCheck': money(net_assets - liab - susp - equity),
            'tbCumulative': {'debit': money(tb_dr), 'credit': money(tb_cr)},
            'openingGross': {'debit': money(sum(open_gdr.values())), 'credit': money(sum(open_gcr.values()))},
        }
        check(net_assets - liab - susp - equity == 0)
        check(tb_dr == tb_cr)
        for a in codes:
            d, c = mov.get((y, a), (0, 0))
            o = opening.get(a, 0)
            cl = o + d - c
            oracle['accounts'].setdefault(a, {})[str(y)] = {
                'openGrossDr': money(open_gdr.get(a, 0)), 'openGrossCr': money(open_gcr.get(a, 0)),
                'openDr': money(max(o, 0)), 'openCr': money(max(-o, 0)), 'movDr': money(d), 'movCr': money(c),
                'closeDr': money(max(cl, 0)), 'closeCr': money(max(-cl, 0))}
        snapshots[y] = dict(cum)
        prior_net_result += net_result
    last = years[-1]
    # cumulative-vs-yearly identity: sum of yearly results = retained + current at the last year (no closing entries)
    oracle['sumOfYearlyResults'] = money(prior_net_result)

    # ------------------------------------------------------------------ 2025 (last year) statements: replay vs cached
    y = last
    m = lambda a: mov.get((y, a), (0, 0))      # noqa: E731
    by_cls = lambda cl: sum(m(a)[1] - m(a)[0] for a in codes if cls_of[a] == cl)   # noqa: E731
    revenue = by_cls('إيرادات') + by_cls('إيرادات مدين')
    cogs = -by_cls('تكلفة مبيعات')
    opex = -by_cls('مصروفات')
    nond = -by_cls('مصروفات غير واجبة الخصم')
    replay_is = {'revenue': revenue, 'cogs': cogs, 'grossProfit': revenue - cogs, 'expenses': opex,
                 'nonDeductible': nond, 'totalExpenses': opex + nond, 'netResult': revenue - cogs - opex - nond}
    cum_last = snapshots[last]
    replay_bs = {
        'assetsGross': sum(cum_last[a] for a in codes if cls_of[a] == 'أصول'),
        'accumDep': -sum(cum_last[a] for a in codes if cls_of[a] == 'مجمع إهلاك'),
    }
    replay_bs['netAssets'] = replay_bs['assetsGross'] - replay_bs['accumDep']
    s5 = cached_by_label(sheets[5], 1, 2)
    s6 = cached_by_label(sheets[6], 1, 2)
    cached_is = {'revenue': s5.get('إيرادات النشاط'), 'cogs': s5.get('تكلفة المبيعات'), 'grossProfit': s5.get('مجمل الربح (الخسارة)'),
                 'expenses': s5.get('مرتبات ومكافآت وأتعاب ومصروفات عمومية'),
                 'nonDeductible': s5.get('مصروفات غير واجبة الخصم ضريبياً'), 'totalExpenses': s5.get('إجمالي المصروفات'),
                 'netResult': s5.get('صافي الربح (الخسارة) عن السنة')}
    cached_bs = {'assetsGross': s6.get('الأصول بالتكلفة'), 'accumDep': s6.get('يخصم: مجمع الإهلاك والإطفاء'),
                 'netAssets': s6.get('إجمالي الأصول'), 'capital': s6.get('رأس المال الصادر') or s6.get('رأس المال المصدر'),
                 'capitalCalled': s6.get('يخصم: رأس المال تحت الطلب'), 'settlementShares': s6.get('حصص المساهمين مقابل عقد التسوية'),
                 'retained': s6.get('أرباح (خسائر) مرحلة'), 'currentResult': s6.get('صافي ربح (خسارة) السنة'),
                 'equity': s6.get('إجمالي حقوق الملكية'), 'liabilities': s6.get('إجمالي الالتزامات'),
                 'liabilitiesAndEquity': s6.get('إجمالي الالتزامات وحقوق الملكية'), 'balanceTest': s6.get('اختبار التوازن')}
    for k, v in replay_is.items():
        cv = cents_cached(cached_is.get(k))
        if cv is not None:
            check(cv == v)
    for k, v in replay_bs.items():
        cv = cents_cached(cached_bs.get(k))
        if cv is not None:
            check(cv == v)
    ylast = oracle['years'][str(last)]
    for k_c, k_o in (('equity', 'equity'), ('liabilities', 'liabilities')):
        cv = cents_cached(cached_bs.get(k_c))
        if cv is not None:
            check(cv == int(round(ylast[k_o] * 100)))
    share_codes = [a for a in codes if a.startswith('32') and cls_of[a] == 'حقوق ملكية']
    retained_codes = [a for a in ('3150', '3160') if a in cum_last]
    equity_other = [a for a in codes if cls_of[a] == 'حقوق ملكية' and a not in share_codes and a not in retained_codes]
    replay_eq = {
        'capital': -sum(cum_last[a] for a in equity_other),
        'capitalCalled': sum(cum_last[a] for a in codes if cls_of[a] == 'حقوق ملكية مدين'),
        'settlementShares': -sum(cum_last[a] for a in share_codes),
        'retained': int(round(ylast['equityParts']['priorResults'] * 100)) - sum(cum_last[a] for a in retained_codes),
        'currentResult': int(round(ylast['equityParts']['currentResult'] * 100)),
    }
    for k, v in replay_eq.items():
        cv = cents_cached(cached_bs.get(k))
        if cv is not None:
            check(cv == v)
    replay_bs.update(replay_eq)
    oracle['reports'] = {str(last): {
        'incomeStatement': {'replay': {k: money(v) for k, v in replay_is.items()},
                            'cached': {k: num_or_none(v) for k, v in cached_is.items()}},
        'balanceSheet': {'replay': {k: money(v) for k, v in replay_bs.items()},
                         'cached': {k: num_or_none(v) if not isinstance(v, str) else v for k, v in cached_bs.items()},
                         'derived': {'liabilities': ylast['liabilities'], 'equity': ylast['equity'],
                                     'suspense': ylast['suspense'], 'priorResults': ylast['equityParts']['priorResults'],
                                     'currentResult': ylast['equityParts']['currentResult']}}}}

    # ------------------------------------------------------------------ sheet 3 (ledger of the selected year) and 4 (TB)
    ws3 = sheets[3]
    s3_year = ws3['B1'].value
    sheet3_rows = 0
    s3_match = None
    if isinstance(s3_year, (int, float)) and int(s3_year) in years:
        sy = int(s3_year)
        s3_match = sy
        for r in range(4, ws3.max_row + 1):
            code = ws3.cell(r, 1).value
            if code is None or not re.fullmatch(r'[1-9]\d{3}', str(code).strip()):
                continue
            code = str(code).strip()
            if code not in oracle['accounts']:
                check(False)
                continue
            sheet3_rows += 1
            rec = oracle['accounts'][code][str(sy)]
            for col, key in ((4, 'openGrossDr'), (5, 'openGrossCr'), (6, 'movDr'), (7, 'movCr'), (8, 'closeDr'), (9, 'closeCr')):
                cv = cents_cached(ws3.cell(r, col).value)
                if cv is not None:
                    check(cv == int(round(rec[key] * 100)))
        for r in range(1, ws3.max_row + 1):
            if clean(ws3.cell(r, 2).value) == 'الإجمالي':
                for col, key in ((4, 'openGrossDr'), (5, 'openGrossCr'), (6, 'movDr'), (7, 'movCr'), (8, 'closeDr'), (9, 'closeCr')):
                    tot = sum(int(round(oracle['accounts'][a][str(sy)][key] * 100)) for a in codes)
                    cv = cents_cached(ws3.cell(r, col).value)
                    if cv is not None:
                        check(cv == tot)
    ws4 = sheets[4]
    tb_rows = 0
    tb_cached = None
    for r in range(1, ws4.max_row + 1):
        if clean(ws4.cell(r, 2).value) == 'الإجمالي':
            tb_cached = {'debit': num_or_none(ws4.cell(r, 3).value), 'credit': num_or_none(ws4.cell(r, 4).value)}
    for r in range(1, ws4.max_row + 1):
        code = ws4.cell(r, 1).value
        if code is not None and re.fullmatch(r'[1-9]\d{3}', str(code).strip()):
            tb_rows += 1
            rec = oracle['accounts'][str(code).strip()][str(s3_match or last)]
            for col, key in ((3, 'closeDr'), (4, 'closeCr')):
                cv = cents_cached(ws4.cell(r, col).value)
                if cv is not None:
                    check(cv == int(round(rec[key] * 100)))
    nz = [a for a in codes if cum_last[a] != 0]
    oracle['reports'][str(last)]['trialBalance'] = {
        'debit': ylast['tbCumulative']['debit'], 'credit': ylast['tbCumulative']['credit'], 'accountsNonZero': len(nz),
        'cached': tb_cached, 'sheet3ReportYear': s3_match, 'sheet3RowsCompared': sheet3_rows, 'sheet4RowsCompared': tb_rows}
    if tb_cached and tb_cached['debit'] is not None:
        check(cents_cached(tb_cached['debit']) == int(round(ylast['tbCumulative']['debit'] * 100)))
        check(cents_cached(tb_cached['credit']) == int(round(ylast['tbCumulative']['credit'] * 100)))

    # ------------------------------------------------------------------ sheet 9: party balances (by account binding)
    ws9 = sheets[9]
    party_cached, party_replay = {}, {}
    for r in range(1, ws9.max_row + 1):
        for code_col, bal_col in ((3, 4), (5, 6)):
            code = ws9.cell(r, code_col).value
            if code is not None and re.fullmatch(r'[1-9]\d{3}', str(code).strip()):
                code = str(code).strip()
                party_cached[code] = num_or_none(ws9.cell(r, bal_col).value)
    for a in FUNDING_ACCOUNTS:
        if a in cum_last:
            party_replay[a] = -cum_last[a]
    for a, v in party_cached.items():
        if v is not None and a in party_replay:
            check(cents_cached(v) == party_replay[a])
    oracle['partyBalances'] = {'asOfYear': last, 'cached': party_cached,
                               'replay': {a: money(v) for a, v in party_replay.items()}}

    # ------------------------------------------------------------------ sheet 20: cash change row
    ws20 = sheets[20]
    cash_codes = [a for a in codes if a.startswith(CASH_PREFIX)]
    cash_replay = {}
    for yy in years:
        cash_replay[str(yy)] = money(sum(mov.get((yy, a), (0, 0))[0] - mov.get((yy, a), (0, 0))[1] for a in cash_codes))
    cash_cached = {}
    hdr_row = None
    for r in range(1, ws20.max_row + 1):
        if clean(ws20.cell(r, 1).value) == 'البند':
            hdr_row = r
    if hdr_row:
        col_year = {}
        for c in range(2, ws20.max_column + 1):
            h = ws20.cell(hdr_row, c).value
            if h is not None and re.fullmatch(r'\d{4}', to_western(str(h)).strip()):
                col_year[c] = to_western(str(h)).strip()
        for r in range(hdr_row + 1, ws20.max_row + 1):
            lab = clean(ws20.cell(r, 1).value)
            if lab.startswith('هـ)'):
                for c, yy in col_year.items():
                    v = ws20.cell(r, c).value
                    cash_cached[yy] = num_or_none(v)
            if lab.startswith('تحقق'):
                oracle['cashCheckRowCached'] = {yy: num_or_none(ws20.cell(r, c).value) for c, yy in col_year.items()}
    for yy, v in cash_cached.items():
        if v is not None and yy in cash_replay:
            check(cents_cached(v) == int(round(cash_replay[yy] * 100)))
    oracle['cashChange'] = {'cached': cash_cached, 'replay': cash_replay}

    # ------------------------------------------------------------------ sheets 15, 16: register totals
    ws15 = sheets[15]
    legal_items = []
    for r in range(1, ws15.max_row + 1):
        n, amt = ws15.cell(r, 1).value, ws15.cell(r, 3).value
        if isinstance(n, (int, float)) and isinstance(amt, (int, float)) and ws15.cell(r, 2).value is not None:
            legal_items.append(int(round(amt * 100)))
    legal_total_cached = None
    for r in range(1, ws15.max_row + 1):
        if clean(ws15.cell(r, 1).value) == 'الإجمالي':
            legal_total_cached = ws15.cell(r, 3).value
    entry_sums = j.groupby('no')['dr'].sum()
    match_entries = [int(n) for n, s in entry_sums.items() if s == sum(legal_items)]
    by_acct = {}
    if match_entries:
        sub = j[(j['no'] == match_entries[0]) & (j['dr'] > 0)].groupby('acct')['dr'].sum()
        by_acct = {a: money(v) for a, v in sub.items()}
    oracle['legalExpenses'] = {'items': len(legal_items), 'sumOfItems': money(sum(legal_items)), 'cachedTotal': legal_total_cached,
                               'entryMatchingTotal': match_entries[0] if match_entries else None, 'entryDebitByAccount': by_acct}
    if legal_total_cached is not None:
        check(cents_cached(legal_total_cached) == sum(legal_items))
    ws16 = sheets[16]
    pay_rows, pay_total_cached = [], None
    for r in range(1, ws16.max_row + 1):
        a = ws16.cell(r, 1).value
        if isinstance(a, (int, float)) and isinstance(ws16.cell(r, 5).value, (int, float)):
            pay_rows.append((int(a), int(round(ws16.cell(r, 5).value * 100))))
        if clean(a) == 'الإجمالي':
            pay_total_cached = ws16.cell(r, 5).value
    oracle['payroll'] = {'years': {str(a): money(v) for a, v in pay_rows}, 'sumOfYears': money(sum(v for _, v in pay_rows)),
                         'cachedTotal': pay_total_cached}
    if pay_total_cached is not None:
        check(cents_cached(pay_total_cached) == sum(v for _, v in pay_rows))

    # ------------------------------------------------------------------ sheet 8 / 12: settlement shares
    oracle['settlementShares'] = {
        'ledgerLast': {a: money(-cum_last[a]) for a in share_codes},
        'ledgerLastTotal': money(-sum(cum_last[a] for a in share_codes)),
        'creditsTotal': money(int(j[j['acct'].isin(share_codes)]['cr'].sum())),
    }
    ws8 = sheets[8]
    for r in range(1, ws8.max_row + 1):
        if clean(ws8.cell(r, 1).value) == 'الإجمالي':
            oracle['settlementShares']['sheet8Total'] = num_or_none(ws8.cell(r, 4).value)

    # ------------------------------------------------------------------ sheet 18: allocation outputs (cached)
    ws18 = sheets[18]
    alloc = {'scenarios': []}
    hdr_rows = [r for r in range(1, ws18.max_row + 1) if clean(ws18.cell(r, 2).value) == 'نسبة الأساس']
    for h in hdr_rows:
        sc = {'rows': [], 'total': None, 'price': None}
        for r in range(h + 1, ws18.max_row + 1):
            lab = clean(ws18.cell(r, 1).value)
            if not lab:
                break
            if lab == 'الإجمالي':
                sc['total'] = {'totalShares': num_or_none(ws18.cell(r, 14).value), 'sharesTo2024': num_or_none(ws18.cell(r, 10).value),
                               'basePct': num_or_none(ws18.cell(r, 2).value), 'pctAtClose': num_or_none(ws18.cell(r, 15).value)}
            elif lab.startswith('سعر'):
                sc['price'] = {'price2023_2024': num_or_none(ws18.cell(r, 7).value), 'price2025': num_or_none(ws18.cell(r, 13).value)}
            elif lab.startswith('قيمة الشركة'):
                sc['impliedValue'] = num_or_none(ws18.cell(r, 2).value)
                break
            else:
                sc['rows'].append({'sharesAtClose': num_or_none(ws18.cell(r, 14).value), 'pctAtClose': num_or_none(ws18.cell(r, 15).value)})
        pct_sum = sum(x['pctAtClose'] for x in sc['rows'] if x['pctAtClose'] is not None)
        sc['pctAtCloseSum'] = pct_sum
        sc['rowCount'] = len(sc['rows'])
        if sc['rows']:
            check(abs(pct_sum - 1.0) < 1e-9)
        if sc['total'] and sc['total']['totalShares'] is not None:
            check(abs(sum(x['sharesAtClose'] or 0 for x in sc['rows']) - sc['total']['totalShares']) < 1e-6)
        alloc['scenarios'].append(sc)
    oracle['allocation'] = alloc

    # ------------------------------------------------------------------ sheet 19: evidence counts and totals
    ws19 = sheets[19]
    ev = {'included': [0, 0], 'after_close': [0, 0], 'excluded': [0, 0]}
    grades = {'A': [0, 0], 'B': [0, 0], 'C': [0, 0], 'unknown': [0, 0]}
    smap = {'نعم': 'included', 'بعد الإقفال': 'after_close', 'مستبعد': 'excluded'}
    for r in range(1, ws19.max_row + 1):
        st = clean(ws19.cell(r, 9).value)
        amt = ws19.cell(r, 3).value
        if st in smap and isinstance(amt, (int, float)) and ws19.cell(r, 2).value is not None:
            k = smap[st]
            ev[k][0] += 1
            ev[k][1] += int(round(amt * 100))
            g = re.match(r'^[\(（]\s*([أاإبج])', clean(ws19.cell(r, 8).value))
            gk = {'أ': 'A', 'ا': 'A', 'إ': 'A', 'ب': 'B', 'ج': 'C'}.get(g.group(1), 'unknown') if g else 'unknown'
            grades[gk][0] += 1
            grades[gk][1] += int(round(amt * 100))
    oracle['evidence'] = {'status': {k: {'count': v[0], 'total': money(v[1])} for k, v in ev.items()},
                          'grades': {k: {'count': v[0], 'total': money(v[1])} for k, v in grades.items()},
                          'rows': sum(v[0] for v in ev.values())}

    # ------------------------------------------------------------------ depreciation actually posted (by year x accumulated account)
    accum_codes = [a for a in codes if cls_of[a] == 'مجمع إهلاك' and a != '1240']
    dep = {}
    for yy in years:
        row = {a: money(mov.get((yy, a), (0, 0))[1] - mov.get((yy, a), (0, 0))[0]) for a in accum_codes}
        total = sum(mov.get((yy, a), (0, 0))[1] - mov.get((yy, a), (0, 0))[0] for a in accum_codes)
        dep[str(yy)] = {'total': money(total), 'byAccum': {a: v for a, v in row.items() if v != 0}}
    oracle['depreciationPosted'] = dep
    ws7 = sheets[7]
    s7_total = None
    s7_rows = 0
    for r in range(1, ws7.max_row + 1):
        a = ws7.cell(r, 1).value
        if a is not None and re.fullmatch(r'[1-9]\d{3}', str(a).strip()) and isinstance(ws7.cell(r, 3).value, (int, float)):
            s7_rows += 1
        if clean(ws7.cell(r, 2).value) == 'الإجمالي':
            s7_total = ws7.cell(r, 3).value
    oracle['assetsSheet'] = {'rows': s7_rows, 'totalCost': s7_total}

    # ------------------------------------------------------------------ counts
    entries = j.groupby('no').size()
    nos = sorted(int(n) for n in entries.index)
    used = set(j['acct'])
    par = j['الطرف'].fillna('').astype(str).str.strip()
    doc = j['رقم المستند'].fillna('').astype(str).str.strip()
    memo = (j['البيان'].fillna('').astype(str) + ' ' + doc + ' ' + par).map(to_western)
    ar = r'(?<![ء-ي])'
    a_lines = int(memo.str.contains(ar + r'افتراض(?:\s+رقم)?\s*\d{1,3}(?!\d)(?!\s*[-–]\s*\d)', regex=True).sum())
    o_lines = int(memo.str.contains(
        ar + r'(?:بند(?:\s+مفتوح)?|(?:قيد\s+)?(?:ال)?استفسار)\s*\d{1,3}(?!\d)(?!\s*[-–]\s*\d)', regex=True).sum())
    ws10, ws13 = sheets[10], sheets[13]
    oi = [r for r in range(1, ws10.max_row + 1) if re.fullmatch(r'\d+', str(ws10.cell(r, 1).value or '').strip())]
    status_counts = Counter(clean(ws10.cell(r, 6).value) for r in oi)
    asm = [r for r in range(1, ws13.max_row + 1) if re.fullmatch(r'\d+', str(ws13.cell(r, 1).value or '').strip())]
    counts = {
        'lines': int(len(j)), 'entries': int(len(nos)), 'entryNoMin': nos[0], 'entryNoMax': nos[-1],
        'entryNoGaps': nos[-1] - nos[0] + 1 - len(nos), 'accounts': len(codes), 'accountsUsed': len(used),
        'accountsUnused': len(codes) - len(used), 'partyRawDistinct': int(par.nunique()),
        'partyPlaceholderLines': int(par.isin(PLACEHOLDER_TEXTS).sum()), 'documentsDistinct': int(doc[doc != ''].nunique()),
        'openItems': len(oi), 'openItemStatus': dict(sorted(status_counts.items())), 'assumptions': len(asm),
        'linesCitingAssumption': a_lines, 'linesCitingOpenItem': o_lines,
        'linesWithInquiryWord': int(memo.str.contains('استفسار').sum()),
        'costCenters': int(j['مركز التكلفة'].nunique()), 'sectors': int(j['القطاع القانوني'].nunique()),
        'totalDebit': money(int(j['dr'].sum())), 'totalCredit': money(int(j['cr'].sum())),
        'assetRowsSheet7': s7_rows,
    }
    check(counts['totalDebit'] == counts['totalCredit'])
    bad_entries = [int(n) for n, g in j.groupby('no') if g['dr'].sum() != g['cr'].sum()]
    check(not bad_entries)
    counts['unbalancedEntries'] = len(bad_entries)
    if seed_dir and os.path.isdir(seed_dir):
        pdir = os.path.join(seed_dir, 'parties')
        adir = os.path.join(seed_dir, 'assets')
        counts['seedDerived'] = {
            'partiesAfterMerge': len([f for f in os.listdir(pdir)]) if os.path.isdir(pdir) else None,
            'assetsRegister': len([f for f in os.listdir(adir)]) if os.path.isdir(adir) else None,
        }
    oracle['counts'] = counts
    ok = bad[0] == 0
    oracle['selfCheck'] = {'pass': ok, 'checks': checks[0], 'mismatches': bad[0], 'failedCheckLines': dict(failed)}
    return oracle, ok, checks[0], bad[0]


if __name__ == '__main__':
    sys.exit(main())

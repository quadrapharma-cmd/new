"""Workbook reader: turns the 21-sheet workbook into plain Python structures (no decisions are made here).

Sheets are addressed by position AND checked by their Arabic-Indic numeric title prefix ("٢. ..."), so a reordered or
renamed file stops the run instead of being misread. The journal is read by account code (rows without one are
skipped), never by position.
"""
import re
from datetime import date, datetime

import openpyxl

from common import (SeedError, arabic_index, cell_text, clean, exact_cents, iso, num, parse_dmy, sha256_file,
                    to_western)

EXPECTED_SHEETS = 21
JOURNAL_HEADERS = {1: 'التاريخ', 3: 'رقم القيد', 4: 'كود الحساب', 6: 'مدين', 7: 'دائن', 8: 'البيان',
                   9: 'رقم المستند', 10: 'مركز التكلفة', 11: 'الطرف', 12: 'القطاع القانوني'}
CHART_HEADERS = {1: 'الكود', 2: 'اسم الحساب', 3: 'التصنيف'}


class Workbook:
    """Holds the value view (cached results) and, lazily, the formula view of the same file."""

    def __init__(self, path):
        self.path = path
        self.sha256 = sha256_file(path)
        self.values = openpyxl.load_workbook(path, data_only=True)
        self._formulas = None
        props = self.values.properties
        self.created = props.created.isoformat() if props.created else None
        self.modified = props.modified.isoformat() if props.modified else None

    @property
    def formulas(self):
        if self._formulas is None:
            self._formulas = openpyxl.load_workbook(self.path, data_only=False)
        return self._formulas

    def sheet(self, idx, formulas=False):
        book = self.formulas if formulas else self.values
        return book.worksheets[idx]

    def title(self, idx):
        return self.values.worksheets[idx].title


def open_workbook(path):
    wb = Workbook(path)
    n = len(wb.values.worksheets)
    if n != EXPECTED_SHEETS:
        raise SeedError(f'unexpected sheet count: {n} (expected {EXPECTED_SHEETS})')
    for i, ws in enumerate(wb.values.worksheets):
        prefix = arabic_index(i) + '.'
        if not ws.title.strip().startswith(prefix):
            raise SeedError(f'sheet {i} title does not start with "{prefix}"')
    return wb


# --------------------------------------------------------------------------- generic helpers
def val(c):
    """Cell value cleaned for storage in JSON registers: dates -> ISO, strings trimmed, numbers kept."""
    v = c.value if hasattr(c, 'value') else c
    if v is None:
        return None
    if isinstance(v, (datetime, date)):
        return iso(v)
    if isinstance(v, str):
        s = clean(v)
        return s if s != '' else None
    if isinstance(v, float) and v == int(v) and abs(v) < 1e15:
        return int(v)
    return v


def row_values(ws, r, c1=1, c2=None):
    c2 = c2 or ws.max_column
    return [val(ws.cell(r, c)) for c in range(c1, c2 + 1)]


def blank_row(ws, r, c1=1, c2=None):
    return all(v is None for v in row_values(ws, r, c1, c2))


def first_text(ws, r, c1=1, c2=None):
    for v in row_values(ws, r, c1, c2):
        if isinstance(v, str):
            return v
    return None


def as_int(v):
    n = num(v)
    if n is None:
        return None
    return int(n) if float(n) == int(n) else None


def find_rows(ws, pred, c1=1, c2=None, r1=1, r2=None):
    r2 = r2 or ws.max_row
    return [r for r in range(r1, r2 + 1) if pred(row_values(ws, r, c1, c2))]


def notes_of(ws, consumed):
    """Every row not consumed by a table that holds text -> a note (heading rows are kept in order)."""
    out = []
    for r in range(1, ws.max_row + 1):
        if r in consumed:
            continue
        vals = [v for v in row_values(ws, r) if v is not None]
        if not vals:
            continue
        texts = [str(v) for v in vals]
        marker = texts[0] if len(texts) > 1 and len(texts[0]) <= 2 else None
        body = ' '.join(texts[1:]) if marker is not None else ' '.join(texts)
        out.append({'row': r, 'marker': marker, 'text': body})
    return out


# --------------------------------------------------------------------------- sheet 0: company identity
def read_company(wb):
    ws = wb.sheet(0)
    title = clean(ws['A1'].value)
    m = re.match(r'^(.*?)\s*[\(（]\s*(.*?)\s*سابق(?:اً|ا|ً)?\s*[\)）]\s*$', title)
    if m:
        name, former = clean(m.group(1)), clean(m.group(2))
    else:
        name, former = title, ''
    info = {}
    id_rows = set()
    for r in range(1, ws.max_row + 1):
        label = clean(ws.cell(r, 1).value)
        if ws.cell(r, 2).value is None or len(label) > 30:
            continue
        key = 'tax' if 'ضريب' in label else 'cr' if 'سجل' in label else 'inc' if 'تأسيس' in label else None
        if key and key not in info:
            info[key] = ws.cell(r, 2).value
            id_rows.add(r)
    inc = info.get('inc')
    incorporated = ''
    if isinstance(inc, (datetime, date)):
        incorporated = iso(inc)
    elif inc:
        incorporated = parse_dmy(str(inc)) or ''
    guide = []
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        if a is not None and b is not None and r not in id_rows:
            guide.append({'row': r, 'marker': cell_text(a), 'text': cell_text(b)})
    return {
        'name': name, 'formerName': former,
        'taxNo': cell_text(info.get('tax')), 'crNo': cell_text(info.get('cr')),
        'incorporated': incorporated, 'guide': guide,
    }


# --------------------------------------------------------------------------- sheet 1: chart of accounts
def read_chart(wb):
    ws = wb.sheet(1)
    for c, label in CHART_HEADERS.items():
        if clean(ws.cell(1, c).value) != label:
            raise SeedError(f'chart header mismatch at column {c}')
    sections, accounts = [], []
    for r in range(2, ws.max_row + 1):
        code, name, cls = ws.cell(r, 1).value, ws.cell(r, 2).value, ws.cell(r, 3).value
        if code is None and name is None:
            continue
        if name is None and cls is None:
            sections.append({'title': clean(code), 'row': r, 'before': len(accounts)})
            continue
        code = cell_text(code)
        if not re.fullmatch(r'[1-9]\d{3}', code):
            raise SeedError(f'chart row {r}: account code is not a 4-digit code')
        accounts.append({'code': code, 'name': clean(name), 'cls': clean(cls), 'row': r})
    seen = set()
    for a in accounts:
        if a['code'] in seen:
            raise SeedError('duplicate account code in chart')
        seen.add(a['code'])
    return {'accounts': accounts, 'sections': sections}


# --------------------------------------------------------------------------- sheet 2: journal
def read_journal(wb):
    ws = wb.sheet(2)
    for c, label in JOURNAL_HEADERS.items():
        if clean(ws.cell(1, c).value) != label:
            raise SeedError(f'journal header mismatch at column {c}')
    lines, skipped = [], 0
    for r in range(2, ws.max_row + 1):
        code = ws.cell(r, 4).value
        if code is None or clean(code) == '':
            if any(ws.cell(r, c).value is not None for c in (6, 7, 8, 9)):
                skipped += 1  # content without an account code: reported, never imported
            continue
        d = ws.cell(r, 1).value
        if isinstance(d, datetime):
            d = d.date()
        if not isinstance(d, date):
            raise SeedError(f'journal row {r}: date is missing or not a date')
        no = as_int(ws.cell(r, 3).value)
        if no is None or no < 1:
            raise SeedError(f'journal row {r}: entry number is missing')
        dr, cr = exact_cents(num(ws.cell(r, 6).value)), exact_cents(num(ws.cell(r, 7).value))
        if dr is None or cr is None:
            raise SeedError(f'journal row {r}: amount has more than two decimals')
        if dr < 0 or cr < 0:
            raise SeedError(f'journal row {r}: negative amount')
        if (dr > 0) == (cr > 0):
            raise SeedError(f'journal row {r}: line must have exactly one side')
        lines.append({
            'row': r, 'date': d, 'no': no, 'acct': cell_text(code), 'dr': dr, 'cr': cr,
            'memo': clean(ws.cell(r, 8).value), 'doc': clean(ws.cell(r, 9).value), 'cc': clean(ws.cell(r, 10).value),
            'party': clean(ws.cell(r, 11).value), 'sector': clean(ws.cell(r, 12).value),
            'status': clean(ws.cell(r, 13).value),
        })
    return {'lines': lines, 'skippedRows': skipped}


# --------------------------------------------------------------------------- sheet 7: assets
def read_assets_sheet(wb):
    ws = wb.sheet(7)
    start = parse_dmy(clean(ws['A2'].value))
    items = []
    consumed = {1, 2, 4}
    for r in range(5, ws.max_row + 1):
        a = ws.cell(r, 1).value
        if a is None or not re.fullmatch(r'[1-9]\d{3}', cell_text(a)):
            if first_text(ws, r) == 'الإجمالي':
                consumed.add(r)
            break
        consumed.add(r)
        items.append({'acct': cell_text(a), 'name': clean(ws.cell(r, 2).value), 'cost': num(ws.cell(r, 3).value),
                      'rate': num(ws.cell(r, 4).value), 'row': r})
    return {'startDate': start, 'items': items, 'consumed': consumed}


# --------------------------------------------------------------------------- generic keyed tables
def read_keyed(ws, header_row, keys, first_row=None, stop_labels=('الإجمالي',), max_col=None):
    """Read rows under a header into dicts. keys: key names for columns 1..n (None = skip column).
    Stops at the first fully blank row, or at a row labelled in stop_labels (returned as `total`)."""
    first_row = first_row or header_row + 1
    width = max_col or len(keys)
    labels = row_values(ws, header_row, 1, width)
    items, total, consumed = [], None, {header_row}
    for r in range(first_row, ws.max_row + 1):
        vals = row_values(ws, r, 1, width)
        if all(v is None for v in vals):
            break
        lab = next((v for v in vals if isinstance(v, str)), None)
        if lab in stop_labels:
            total = {k: v for k, v in zip(keys, vals) if k}
            consumed.add(r)
            break
        consumed.add(r)
        items.append({k: v for k, v in zip(keys, vals) if k})
    columns = [{'key': k, 'label': l} for k, l in zip(keys, labels) if k]
    return {'items': items, 'total': total, 'columns': columns, 'consumed': consumed}


_EMPTY_TABLE = {'items': [], 'total': None, 'columns': [], 'consumed': set()}


# --------------------------------------------------------------------------- sheets 8, 9
def read_settlement(wb):
    ws = wb.sheet(8)
    hdr = find_rows(ws, lambda v: v[0] == 'المساهم')
    if not hdr:
        return dict(_EMPTY_TABLE)
    return read_keyed(ws, hdr[0], ['holder', 'legacyCode', 'pct', 'credit', 'paid', 'balance'])


def read_partner_balances(wb):
    ws = wb.sheet(9)
    hdr = find_rows(ws, lambda v: v[0] == 'الطرف')
    if not hdr:
        return dict(_EMPTY_TABLE)
    return read_keyed(ws, hdr[0], ['party', 'capacity', 'fundingAcct', 'funding', 'accrualAcct', 'accrual', 'total'])


# --------------------------------------------------------------------------- sheets 10, 13: registers
def read_open_items(wb):
    ws = wb.sheet(10)
    hdr = find_rows(ws, lambda v: v[0] == '#')
    if not hdr:
        raise SeedError('open items header not found')
    items = []
    for r in range(hdr[0] + 1, ws.max_row + 1):
        n = as_int(ws.cell(r, 1).value)
        if n is None:
            break
        items.append({
            'n': n, 'yearRaw': val(ws.cell(r, 2)), 'item': clean(ws.cell(r, 3).value),
            'effect': clean(ws.cell(r, 4).value), 'docRequired': clean(ws.cell(r, 5).value),
            'statusText': clean(ws.cell(r, 6).value), 'row': r,
        })
    return {'items': items}


def read_assumptions(wb):
    ws = wb.sheet(13)
    hdr = find_rows(ws, lambda v: v[0] == '#')
    if not hdr:
        raise SeedError('assumptions header not found')
    items = []
    for r in range(hdr[0] + 1, ws.max_row + 1):
        n = as_int(ws.cell(r, 1).value)
        if n is None:
            break
        items.append({'n': n, 'text': clean(ws.cell(r, 2).value), 'effect': clean(ws.cell(r, 3).value),
                      'docRequired': clean(ws.cell(r, 4).value), 'ifNotReceived': clean(ws.cell(r, 5).value)})
    return {'items': items}


# --------------------------------------------------------------------------- sheet 11: close status
def read_close_status(wb):
    ws = wb.sheet(11)
    years, general, needed, status_rows = [], [], [], []
    section = None
    consumed = set()
    for r in range(1, ws.max_row + 1):
        a, b = ws.cell(r, 1).value, ws.cell(r, 2).value
        at = clean(a)
        if at == 'السنة' and clean(b) == 'حالة الإقفال':
            section = 'main'
            consumed.add(r)
            continue
        if at.startswith('تحفظ عام على'):
            section = 'general'
            consumed.add(r)
            general.append({'row': r, 'text': at})
            continue
        if at.startswith('ما يلزم'):
            section = 'needed'
            consumed.add(r)
            continue
        if section == 'main' and re.fullmatch(r'\d{4}', to_western(at)):
            consumed.add(r)
            last = ws.cell(r, 6).value
            years.append({
                'year': int(to_western(at)), 'statusText': clean(b), 'sources': clean(ws.cell(r, 3).value),
                'notRecorded': clean(ws.cell(r, 4).value), 'text': clean(ws.cell(r, 5).value),
                'lastUpdate': (iso(last) if isinstance(last, (datetime, date)) else parse_dmy(str(last or ''))),
                'row': r,
            })
        elif section == 'general':
            consumed.add(r)
            if b is not None:
                general.append({'row': r, 'text': clean(b)})
        elif section == 'needed':
            consumed.add(r)
            if at == '•':
                general.append({'row': r, 'text': clean(b)})
                section = 'closing'
            elif b is not None:
                needed.append({'row': r, 'scope': at, 'text': clean(b)})
        elif section == 'closing':
            consumed.add(r)
            if at == '•' and b is not None:
                general.append({'row': r, 'text': clean(b)})
            elif re.fullmatch(r'\d{4}', to_western(at)) and b is not None:
                status_rows.append({'year': int(to_western(at)), 'text': clean(b), 'row': r})
    return {'years': years, 'general': general, 'needed': needed, 'statusRows': status_rows, 'consumed': consumed}


# --------------------------------------------------------------------------- sheets 12, 14, 15, 16
def read_contingent(wb):
    ws = wb.sheet(12)
    hdr = find_rows(ws, lambda v: v[0] == 'الاسم' and v[1] == 'الصفة')
    if not hdr:
        return dict(_EMPTY_TABLE)
    return read_keyed(ws, hdr[0], ['name', 'role', 'entitlement', 'booked', 'difference'])


def read_financier_uses(wb):
    ws = wb.sheet(14)
    hdr = find_rows(ws, lambda v: v[0] == 'البيان' and v[1] == 'السنة')
    cols = [{'key': 'label', 'label': 'البيان'}, {'key': 'year', 'label': 'السنة'},
            {'key': 'amount', 'label': 'المبلغ'}, {'key': 'counterAccount', 'label': 'الحساب المقابل'}]
    if not hdr:
        return {'items': [], 'columns': cols, 'consumed': set()}
    h = hdr[0]
    items, consumed = [], {h}
    for r in range(h + 1, ws.max_row + 1):
        vals = row_values(ws, r, 1, 4)
        if first_text(ws, r) == 'ملاحظات':
            break
        if all(v is None for v in vals):
            continue
        consumed.add(r)
        if vals[1] is None and vals[2] is None:
            items.append({'label': vals[0], 'year': None, 'amount': None, 'counterAccount': None, 'heading': True})
        else:
            items.append({'label': vals[0], 'year': vals[1], 'amount': num(vals[2]), 'counterAccount': vals[3],
                          'heading': False})
    return {'items': items, 'columns': cols, 'consumed': consumed}


def read_legal(wb):
    ws = wb.sheet(15)
    hdr = find_rows(ws, lambda v: v[0] == '#' and v[1] == 'البند')
    if not hdr:
        return {'items': [], 'summary': [], 'total': None, 'consumed': set()}
    h = hdr[0]
    items, consumed = [], {h}
    r = h + 1
    while r <= ws.max_row:
        n = as_int(ws.cell(r, 1).value)
        if n is None:
            break
        consumed.add(r)
        items.append({'n': n, 'item': clean(ws.cell(r, 2).value), 'amount': num(ws.cell(r, 3).value),
                      'category': clean(ws.cell(r, 4).value), 'actualYear': val(ws.cell(r, 5))})
        r += 1
    summary, total = [], None
    marker = find_rows(ws, lambda v: v[0] == 'ملخص التصنيف')
    if marker:
        consumed.add(marker[0])
        for rr in range(marker[0] + 1, ws.max_row + 1):
            if blank_row(ws, rr):
                break
            consumed.add(rr)
            row = {'label': clean(ws.cell(rr, 1).value), 'amount': num(ws.cell(rr, 3).value)}
            if row['label'] == 'الإجمالي':
                total = row['amount']
            else:
                summary.append(row)
    return {'items': items, 'summary': summary, 'total': total, 'consumed': consumed}


def read_payroll(wb):
    ws = wb.sheet(16)
    hdr = find_rows(ws, lambda v: v[0] == 'السنة')
    if not hdr:
        return {'items': [], 'columns': [], 'groupLabels': [], 'total': None, 'consumed': set()}
    h = hdr[0]
    labels = row_values(ws, h, 1, 7)
    items, consumed, total = [], {h}, None
    for r in range(h + 1, ws.max_row + 1):
        vals = row_values(ws, r, 1, 7)
        if all(v is None for v in vals):
            break
        if vals[0] == 'الإجمالي':
            consumed.add(r)
            total = {'groups': [num(v) for v in vals[1:4]], 'total': num(vals[4])}
            break
        y = as_int(vals[0])
        if y is None:
            break
        consumed.add(r)
        items.append({'year': y, 'groups': [num(v) for v in vals[1:4]], 'total': num(vals[4]),
                      'taxStatus': vals[5], 'insuranceStatus': vals[6]})
    return {'items': items, 'groupLabels': labels[1:4], 'total': total,
            'columns': [{'key': 'year', 'label': labels[0]}, {'key': 'groups', 'label': 'بنود المرتبات'},
                        {'key': 'total', 'label': labels[4]}, {'key': 'taxStatus', 'label': labels[5]},
                        {'key': 'insuranceStatus', 'label': labels[6]}],
            'consumed': consumed}


# --------------------------------------------------------------------------- sheet 17: accountant reconciliation
def read_reconciliation(wb):
    ws = wb.sheet(17)
    starts = []
    for r in range(1, ws.max_row + 1):
        t = to_western(clean(ws.cell(r, 1).value))
        m = re.match(r'^\((\d+)\)\s*(.*)$', t)
        if m and ws.cell(r, 2).value is None:
            starts.append((r, int(m.group(1)), clean(ws.cell(r, 1).value)))
    sections, consumed = [], set()
    for i, (r, idx, title) in enumerate(starts):
        end = starts[i + 1][0] - 1 if i + 1 < len(starts) else ws.max_row
        consumed.add(r)
        hdr = r + 1
        labels = row_values(ws, hdr, 1, 6)
        width = max((j + 1 for j, v in enumerate(labels) if v is not None), default=0)
        consumed.add(hdr)
        rows = []
        for rr in range(hdr + 1, end + 1):
            vals = row_values(ws, rr, 1, width)
            if all(v is None for v in vals):
                continue
            consumed.add(rr)
            rows.append({'row': rr, 'cells': vals})
        sections.append({'id': idx, 'title': title, 'columns': labels[:width], 'rows': rows})
    return {'sections': sections, 'consumed': consumed}


# --------------------------------------------------------------------------- sheet 18: allocation
_SUM_LITERALS = re.compile(r'^=\s*[\d.]+(?:\s*\+\s*[\d.]+)*\s*$')


def _formula_info(fws, cell):
    f = fws[cell].value
    if not isinstance(f, str) or not f.startswith('='):
        return None
    if _SUM_LITERALS.match(f):
        return {'components': [num(x) for x in re.findall(r'[\d.]+', f)]}
    return {'formula': f}


def read_allocation(wb):
    ws, fws = wb.sheet(18), wb.sheet(18, formulas=True)
    consumed = set()
    inputs = []
    for r in range(1, ws.max_row + 1):
        if re.search('مدخلات', clean(ws.cell(r, 1).value)) and ws.cell(r, 2).value is None:
            consumed.add(r)
            rr = r + 1
            while (rr <= ws.max_row and ws.cell(rr, 1).value is not None
                   and not clean(ws.cell(rr, 1).value).startswith('السيناريو')):
                inputs.append({'label': clean(ws.cell(rr, 1).value), 'value': val(ws.cell(rr, 2)),
                               'note': val(ws.cell(rr, 3)), 'row': rr})
                consumed.add(rr)
                rr += 1
            break
    keys = ['person', 'basePct', 'baseShares', 'inj2022', 'shares2022', 'inj2023', 'shares2023', 'inj2024', 'shares2024',
            'sharesTo2024', 'pctTo2024', 'inj2025', 'shares2025', 'sharesAtClose', 'pctAtClose']
    scen_rows = find_rows(ws, lambda v: v[1] == 'نسبة الأساس')
    scenarios = []
    for si, h in enumerate(scen_rows):
        consumed.add(h)
        title = ''
        for tr in (h - 1, h - 2):
            if tr >= 1 and clean(ws.cell(tr, 1).value).startswith('السيناريو'):
                title = clean(ws.cell(tr, 1).value)
                consumed.add(tr)
                break
        rows, total, price = [], None, None
        for r in range(h + 1, ws.max_row + 1):
            vals = row_values(ws, r, 1, 15)
            lab = vals[0]
            if lab is None:
                break
            if isinstance(lab, str) and lab.startswith('سعر'):
                price = dict(price or {}, price2023_2024=num(vals[6]), price2025=num(vals[12]))
                consumed.add(r)
                continue
            if isinstance(lab, str) and lab.startswith('قيمة الشركة'):
                price = dict(price or {}, impliedValue=num(vals[1]))
                consumed.add(r)
                break
            consumed.add(r)
            rec = {k: (num(v) if k != 'person' else v) for k, v in zip(keys, vals)}
            rec['row'] = r
            fi = {}
            for col, key in (('D', 'inj2022'), ('F', 'inj2023'), ('H', 'inj2024'), ('L', 'inj2025')):
                info = _formula_info(fws, f'{col}{r}')
                if info:
                    fi[key] = info
            if fi:
                rec['inputFormulas'] = fi
            if lab == 'الإجمالي':
                total = rec
            else:
                rows.append(rec)
        scenarios.append({'id': 'AB'[si] if si < 2 else str(si + 1), 'title': title, 'headerRow': h, 'rows': rows,
                          'total': total, 'price': price})
    deposits, deposit_title = [], ''
    dep_hdr = find_rows(ws, lambda v: v[0] == 'التاريخ' and v[1] == 'المبلغ' and v[3] == 'محتسب؟')
    if dep_hdr:
        h = dep_hdr[0]
        consumed.add(h)
        if h > 1 and ws.cell(h - 1, 1).value is not None:
            deposit_title = clean(ws.cell(h - 1, 1).value)
            consumed.add(h - 1)
        for r in range(h + 1, ws.max_row + 1):
            vals = row_values(ws, r, 1, 4)
            if all(v is None for v in vals):
                break
            consumed.add(r)
            deposits.append({'date': vals[0], 'amount': num(vals[1]), 'text': vals[2], 'counted': vals[3]})
    return {'inputs': inputs, 'scenarios': scenarios, 'deposits': deposits, 'depositTitle': deposit_title,
            'consumed': consumed}


# --------------------------------------------------------------------------- sheet 19: funding evidence
_GRADE = {'أ': 'A', 'ب': 'B', 'ج': 'C'}


def normalise_grade(text):
    """'(ب) + تطابق مع المحاسب' -> ('B', '+ تطابق مع المحاسب'); unknown -> (None, text)."""
    t = clean(text)
    m = re.match(r'^[\(（]\s*([أاإبج])\s*[\)）]\s*(.*)$', t)
    if not m:
        return None, t
    ch = 'أ' if m.group(1) in 'اإ' else m.group(1)
    return _GRADE.get(ch), clean(m.group(2))


def read_evidence(wb):
    ws = wb.sheet(19)
    hdr = find_rows(ws, lambda v: v[0] == 'الشريك' and v[1] == 'التاريخ')
    if not hdr:
        return {'items': [], 'summary': [], 'consumed': set()}
    h = hdr[0]
    consumed = {h}
    items, section = [], None
    summary, r_end = [], ws.max_row + 1
    for r in range(h + 1, ws.max_row + 1):
        vals = row_values(ws, r, 1, 9)
        if all(v is None for v in vals):
            continue
        if isinstance(vals[0], str) and vals[0].startswith('ملخص'):
            r_end = r
            break
        consumed.add(r)
        if vals[1] is None and all(v is None for v in vals[2:]):
            section = vals[0]
            continue
        grade, grade_note = normalise_grade(vals[7] or '')
        status_text = vals[8]
        status = {'نعم': 'included', 'بعد الإقفال': 'after_close', 'مستبعد': 'excluded'}.get(status_text, 'unknown')
        items.append({'n': len(items) + 1, 'partner': vals[0], 'section': section, 'date': vals[1], 'amount': num(vals[2]),
                      'type': vals[3], 'document': vals[4], 'location': vals[5], 'confirmation': vals[6],
                      'grade': grade, 'gradeNote': grade_note, 'status': status, 'statusText': status_text, 'row': r})
    if r_end <= ws.max_row:
        consumed.add(r_end)
        for r in range(r_end + 1, ws.max_row + 1):
            vals = row_values(ws, r, 1, 4)
            if all(v is None for v in vals):
                continue
            consumed.add(r)
            if vals[0] == 'الشريك':
                continue
            summary.append({'partner': vals[0], 'included': num(vals[1]), 'afterClose': num(vals[2]), 'total': num(vals[3])})
    return {'items': items, 'summary': summary, 'consumed': consumed}


# --------------------------------------------------------------------------- sheet 20: sources and uses
def read_sources_uses(wb):
    ws = wb.sheet(20)
    hdr = find_rows(ws, lambda v: v[0] == 'البند')
    if not hdr:
        return {'items': [], 'columns': [], 'consumed': set()}
    h = hdr[0]
    labels = [cell_text(v) for v in row_values(ws, h, 1, 9)]
    columns = [{'key': f'c{i}', 'label': l} for i, l in enumerate(labels) if i >= 1 and l]
    col_index = {c['key']: int(c['key'][1:]) for c in columns}
    items, consumed = [], {h}
    for r in range(h + 1, ws.max_row + 1):
        vals = row_values(ws, r, 1, 9)
        lab = vals[0]
        if lab is None:
            continue
        if isinstance(lab, str) and lab.startswith('قراءة'):
            break
        consumed.add(r)
        kind = 'heading' if all(v is None for v in vals[1:]) else 'row'
        items.append({'label': lab, 'kind': kind, 'row': r,
                      'values': {k: num(vals[i]) for k, i in col_index.items() if vals[i] is not None}})
    return {'items': items, 'columns': columns, 'consumed': consumed}


# --------------------------------------------------------------------------- cached statement figures
def read_cached_year(wb):
    """The report year selected in the workbook (sheet 3, cell B1)."""
    return as_int(wb.sheet(3)['B1'].value)
